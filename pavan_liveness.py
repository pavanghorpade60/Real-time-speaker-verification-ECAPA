import os
import warnings
import logging
import glob
import sys
import time
import io
import contextlib
import subprocess
import random
import re
from datetime import datetime
from collections import deque
from difflib import SequenceMatcher

# -------------------------------------------------
# CLEAN CONSOLE / WARNINGS
# -------------------------------------------------
os.environ["TRANSFORMERS_NO_ADVISORY_WARNINGS"] = "1"
os.environ["TOKENIZERS_PARALLELISM"] = "false"
os.environ["HF_HUB_DISABLE_SYMLINKS"] = "1"
os.environ["PYTHONWARNINGS"] = "ignore"
os.environ["TORCH_CPP_LOG_LEVEL"] = "ERROR"
os.environ["TF_CPP_MIN_LOG_LEVEL"] = "3"

warnings.simplefilter("ignore")
warnings.filterwarnings("ignore", category=UserWarning)
warnings.filterwarnings("ignore", category=FutureWarning)
warnings.filterwarnings("ignore", message=".*torchaudio.*")
warnings.filterwarnings("ignore", message=".*torchvision is not available.*")
warnings.filterwarnings("ignore", message=".*The torchaudio backend is switched to.*")
warnings.filterwarnings("ignore", module="speechbrain")
warnings.filterwarnings("ignore", module="torchaudio")
warnings.filterwarnings("ignore", module="speechbrain.utils.torch_audio_backend")

logging.getLogger().setLevel(logging.ERROR)
logging.getLogger("transformers").setLevel(logging.ERROR)
logging.getLogger("speechbrain").setLevel(logging.ERROR)
logging.getLogger("torchaudio").setLevel(logging.ERROR)

import torch
import sounddevice as sd
import numpy as np
import librosa
import pandas as pd
import soundfile as sf

from transformers import pipeline
from openpyxl import load_workbook
from openpyxl.styles import Alignment

with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
    from speechbrain.pretrained import EncoderClassifier

# -------------------------------------------------
# BASE PATHS
# -------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGS_DIR = os.path.join(BASE_DIR, "logs")
RECORDINGS_DIR = os.path.join(BASE_DIR, "recordings")
PAVAN_RECORDINGS_DIR = os.path.join(RECORDINGS_DIR, "pavan")
ECAPA_CACHE_DIR = os.path.join(BASE_DIR, "ecapa_cache")

os.makedirs(LOGS_DIR, exist_ok=True)
os.makedirs(RECORDINGS_DIR, exist_ok=True)
os.makedirs(PAVAN_RECORDINGS_DIR, exist_ok=True)

# -------------------------------------------------
# SETTINGS
# -------------------------------------------------
sd.default.device = 1   # change if needed

SAMPLE_RATE = 16000
BLOCK_SIZE = 1024
DEVICE = "cuda" if torch.cuda.is_available() else "cpu"

# Speaker verification thresholds
BEST_SCORE_THRESHOLD = 0.58
AVG_SCORE_THRESHOLD = 0.44

# VAD thresholds
ENERGY_THRESHOLD = 0.004
SILENCE_LIMIT = 1.8
MIN_SPEECH_SECONDS = 2.8
PRE_SPEECH_BUFFER = 8
MAX_RECORD_SECONDS = 12

# Strict prompt verification
PROMPT_MATCH_THRESHOLD = 0.90
MIN_PROMPT_WORD_COVERAGE = 1.00

# Basic anti-playback / liveness heuristic thresholds
MIN_DYNAMIC_RANGE_DB = 18.0
MAX_CLIPPING_RATIO = 0.025
MIN_ZERO_CROSS_RATE = 0.02
MAX_ZERO_CROSS_RATE = 0.18
MAX_SPECTRAL_FLATNESS = 0.32
MIN_ONSET_STRENGTH = 0.0035
MAX_REPEATED_FRAME_CORR = 0.985

# -------------------------------------------------
# RANDOM PROMPT BANK
# -------------------------------------------------
PROMPT_BANK = [
    "My name is Pavan and this is a live secure access test",
    "Today I am reading this prompt for voice authentication",
    "Pavan is speaking live and requesting secure verification",
    "This is a random prompt and I am speaking it live now",
    "I confirm that this is my live voice for authentication",
    "Secure voice access is being tested by Pavan right now",
    "This sentence is spoken live for identity verification",
    "I am Pavan and I am reading the challenge prompt now",
    "This is a fresh prompt and I am speaking it clearly",
    "Live voice verification is being performed by Pavan"
]

# -------------------------------------------------
# HELPER FUNCTIONS
# -------------------------------------------------
def safe_close_libreoffice():
    try:
        subprocess.run(
            ["taskkill", "/f", "/im", "soffice.bin"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=False
        )
        time.sleep(1)
    except Exception:
        pass


def normalize_text(text):
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def text_similarity(a, b):
    a = normalize_text(a)
    b = normalize_text(b)
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def prompt_word_coverage(expected_prompt, spoken_text):
    expected_words = normalize_text(expected_prompt).split()
    spoken_words = normalize_text(spoken_text).split()

    if not expected_words:
        return 0.0, [], []

    spoken_word_set = set(spoken_words)

    matched_words = [w for w in expected_words if w in spoken_word_set]
    missing_words = [w for w in expected_words if w not in spoken_word_set]

    coverage = len(matched_words) / len(expected_words)
    return round(coverage, 4), matched_words, missing_words


def verify_prompt(expected_prompt, spoken_text):
    similarity = round(text_similarity(expected_prompt, spoken_text), 4)
    coverage, matched_words, missing_words = prompt_word_coverage(expected_prompt, spoken_text)

    prompt_ok = (
        similarity >= PROMPT_MATCH_THRESHOLD and
        coverage >= MIN_PROMPT_WORD_COVERAGE
    )

    return {
        "prompt_ok": prompt_ok,
        "similarity": similarity,
        "coverage": coverage,
        "matched_words": matched_words,
        "missing_words": missing_words
    }


def format_sheet(ws, max_width=40):
    for column in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max_len + 2, max_width)

    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True)


def get_next_run_sheet_name(excel_file):
    if not os.path.exists(excel_file):
        return "Run_1"

    wb = load_workbook(excel_file)
    run_numbers = []

    for name in wb.sheetnames:
        if name.startswith("Run_"):
            try:
                run_numbers.append(int(name.split("_")[1]))
            except Exception:
                pass

    next_run = max(run_numbers) + 1 if run_numbers else 1
    return f"Run_{next_run}"


def repair_existing_audio_links(excel_file, base_dir):
    if not os.path.exists(excel_file):
        return

    recordings_dir = os.path.join(base_dir, "recordings")
    pavan_dir = os.path.join(recordings_dir, "pavan")
    wb = load_workbook(excel_file)

    for sheet_name in wb.sheetnames:
        if not (sheet_name.startswith("Run_") or sheet_name == "MASTER_LOG"):
            continue

        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue

        input_col = None
        match_col = None

        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col).value
            if header == "Input_Audio":
                input_col = col
            elif header == "Matched_Audio":
                match_col = col

        if input_col is None and match_col is None:
            continue

        for row in range(2, ws.max_row + 1):
            if input_col is not None:
                input_cell = ws.cell(row=row, column=input_col)

                input_target = None
                if input_cell.hyperlink and input_cell.hyperlink.target:
                    input_target = str(input_cell.hyperlink.target)
                elif input_cell.value:
                    input_target = str(input_cell.value)

                if input_target:
                    input_name = os.path.basename(input_target.replace("\\", "/"))
                    if input_name.lower().endswith(".wav"):
                        fixed_input_path = os.path.join(recordings_dir, input_name)
                        if os.path.exists(fixed_input_path):
                            input_cell.value = "▶ Play Input"
                            input_cell.hyperlink = os.path.abspath(fixed_input_path)
                            input_cell.style = "Hyperlink"

            if match_col is not None:
                match_cell = ws.cell(row=row, column=match_col)

                match_target = None
                if match_cell.hyperlink and match_cell.hyperlink.target:
                    match_target = str(match_cell.hyperlink.target)
                elif match_cell.value:
                    match_target = str(match_cell.value)

                if match_target:
                    match_name = os.path.basename(match_target.replace("\\", "/"))
                    if match_name.lower().endswith(".wav"):
                        fixed_match_path = os.path.join(pavan_dir, match_name)
                        if os.path.exists(fixed_match_path):
                            match_cell.value = "▶ Play Match"
                            match_cell.hyperlink = os.path.abspath(fixed_match_path)
                            match_cell.style = "Hyperlink"

    wb.save(excel_file)


def rebuild_master_log(excel_file):
    if not os.path.exists(excel_file):
        return

    wb = load_workbook(excel_file)

    columns_order = [
        "Timestamp",
        "Expected_Prompt",
        "Spoken_Text",
        "Prompt_Match",
        "Prompt_Word_Coverage",
        "Missing_Prompt_Words",
        "Detected_Speaker",
        "Input_Audio",
        "Matched_Audio",
        "Best_Matching_Sample",
        "Best_Similarity",
        "Average_Similarity",
        "Pitch_Difference",
        "Spectral_Centroid_Diff",
        "Bandwidth_Difference",
        "Energy_Difference",
        "MFCC_Distance",
        "Liveness_Label",
        "Liveness_Score",
        "Decision"
    ]

    all_rows = []

    for sheet in wb.sheetnames:
        if not sheet.startswith("Run_"):
            continue

        ws = wb[sheet]
        if ws.max_row < 2:
            continue

        row_data = {}
        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx)

            if header in ["Input_Audio", "Matched_Audio"]:
                target = None
                if cell.hyperlink and cell.hyperlink.target:
                    target = str(cell.hyperlink.target)
                elif cell.value:
                    target = str(cell.value)
                row_data[header] = target if target else ""
            else:
                row_data[header] = cell.value

        all_rows.append(row_data)

    if not all_rows:
        return

    master_df = pd.DataFrame(all_rows)
    master_df = master_df.reindex(columns=columns_order)

    if "Timestamp" in master_df.columns:
        master_df["Timestamp"] = master_df["Timestamp"].astype(str)
        master_df = master_df.sort_values(by="Timestamp", ascending=False)

    if "MASTER_LOG" in wb.sheetnames:
        del wb["MASTER_LOG"]
        wb.save(excel_file)

    with pd.ExcelWriter(excel_file, engine="openpyxl", mode="a") as writer:
        master_df.to_excel(writer, sheet_name="MASTER_LOG", index=False)

    wb = load_workbook(excel_file)
    ws = wb["MASTER_LOG"]
    format_sheet(ws, max_width=42)

    input_col = None
    match_col = None

    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header == "Input_Audio":
            input_col = col
        elif header == "Matched_Audio":
            match_col = col

    for row in range(2, ws.max_row + 1):
        if input_col is not None:
            cell = ws.cell(row=row, column=input_col)
            path = cell.value
            if path and str(path).strip() and os.path.exists(str(path)):
                cell.value = "▶ Play Input"
                cell.hyperlink = os.path.abspath(str(path))
                cell.style = "Hyperlink"

        if match_col is not None:
            cell = ws.cell(row=row, column=match_col)
            path = cell.value
            if path and str(path).strip() and os.path.exists(str(path)):
                cell.value = "▶ Play Match"
                cell.hyperlink = os.path.abspath(str(path))
                cell.style = "Hyperlink"

    wb.save(excel_file)


def capture_speech():
    recording = []
    pre_buffer = deque(maxlen=PRE_SPEECH_BUFFER)

    speaking = False
    silence_start = None
    speech_start = None
    total_start = time.time()

    stream = sd.InputStream(
        samplerate=SAMPLE_RATE,
        channels=1,
        blocksize=BLOCK_SIZE
    )

    stream.start()

    try:
        while True:
            audio_chunk, _ = stream.read(BLOCK_SIZE)
            chunk = audio_chunk.flatten()
            energy = np.mean(np.abs(chunk))
            pre_buffer.append(chunk)

            if energy > ENERGY_THRESHOLD:
                if not speaking:
                    speaking = True
                    speech_start = time.time()
                    print("\nSpeech detected... Start reading the prompt.")
                    for buffered in pre_buffer:
                        recording.append(buffered)

                recording.append(chunk)
                silence_start = None

            else:
                if speaking:
                    recording.append(chunk)

                    if silence_start is None:
                        silence_start = time.time()
                    elif time.time() - silence_start > SILENCE_LIMIT:
                        duration = time.time() - speech_start
                        if duration >= MIN_SPEECH_SECONDS:
                            print("Speech ended.")
                            break
                        else:
                            recording = []
                            speaking = False
                            silence_start = None
                            speech_start = None
                            pre_buffer.clear()
                            print("Speech too short. Waiting again...")

            if time.time() - total_start > MAX_RECORD_SECONDS:
                if speaking and len(recording) > 0:
                    print("Max recording window reached. Finalizing captured speech.")
                    break

    finally:
        stream.stop()
        stream.close()

    return recording


def prepare_audio(recording):
    if len(recording) == 0:
        return None

    audio_np = np.concatenate(recording)
    audio_np, _ = librosa.effects.trim(audio_np, top_db=20)

    if len(audio_np) == 0:
        return None

    peak = np.max(np.abs(audio_np))
    if peak > 0:
        audio_np = audio_np / peak

    return audio_np.astype(np.float32)


def extract_embedding(classifier, audio_np):
    waveform = torch.tensor(audio_np).unsqueeze(0).to(DEVICE)

    with torch.no_grad():
        emb = classifier.encode_batch(waveform)
        emb = torch.nn.functional.normalize(emb.squeeze(1), dim=1)

    return emb


def compute_reference_scores(classifier, input_embedding, audio_np):
    audio_files = sorted(glob.glob(os.path.join(PAVAN_RECORDINGS_DIR, "pavan_*.wav")))

    if not audio_files:
        print("No reference samples found in:", PAVAN_RECORDINGS_DIR)
        sys.exit()

    best_score = -1.0
    best_sample = None
    similarities = []

    pitch_diffs = []
    centroid_diffs = []
    bandwidth_diffs = []
    energy_diffs = []
    mfcc_diffs = []

    pitch_input = np.nanmean(librosa.yin(audio_np, fmin=50, fmax=300))
    centroid_input = np.mean(librosa.feature.spectral_centroid(y=audio_np, sr=SAMPLE_RATE))
    bandwidth_input = np.mean(librosa.feature.spectral_bandwidth(y=audio_np, sr=SAMPLE_RATE))
    energy_input = np.mean(librosa.feature.rms(y=audio_np))
    mfcc_input = np.mean(librosa.feature.mfcc(y=audio_np, sr=SAMPLE_RATE, n_mfcc=13), axis=1)

    for file in audio_files:
        sample_audio, _ = librosa.load(file, sr=SAMPLE_RATE)
        sample_audio, _ = librosa.effects.trim(sample_audio, top_db=20)

        peak = np.max(np.abs(sample_audio))
        if peak > 0:
            sample_audio = sample_audio / peak

        sample_audio = sample_audio.astype(np.float32)
        sample_embedding = extract_embedding(classifier, sample_audio)

        score = torch.nn.functional.cosine_similarity(input_embedding, sample_embedding).item()
        similarities.append(score)

        if score > best_score:
            best_score = score
            best_sample = os.path.basename(file)

        pitch = np.nanmean(librosa.yin(sample_audio, fmin=50, fmax=300))
        centroid = np.mean(librosa.feature.spectral_centroid(y=sample_audio, sr=SAMPLE_RATE))
        bandwidth = np.mean(librosa.feature.spectral_bandwidth(y=sample_audio, sr=SAMPLE_RATE))
        energy = np.mean(librosa.feature.rms(y=sample_audio))
        mfcc = np.mean(librosa.feature.mfcc(y=sample_audio, sr=SAMPLE_RATE, n_mfcc=13), axis=1)

        pitch_diffs.append(abs(pitch_input - pitch))
        centroid_diffs.append(abs(centroid_input - centroid))
        bandwidth_diffs.append(abs(bandwidth_input - bandwidth))
        energy_diffs.append(abs(energy_input - energy))
        mfcc_diffs.append(np.linalg.norm(mfcc_input - mfcc))

    avg_similarity = float(np.mean(similarities))

    return {
        "best_score": best_score,
        "best_sample": best_sample,
        "avg_similarity": avg_similarity,
        "pitch_diff": round(float(np.mean(pitch_diffs)), 2),
        "centroid_diff": round(float(np.mean(centroid_diffs)), 2),
        "bandwidth_diff": round(float(np.mean(bandwidth_diffs)), 2),
        "energy_diff": round(float(np.mean(energy_diffs)), 4),
        "mfcc_diff": round(float(np.mean(mfcc_diffs)), 4),
    }


def estimate_liveness(audio_np):
    rms = librosa.feature.rms(y=audio_np, frame_length=1024, hop_length=256)[0]
    zcr = librosa.feature.zero_crossing_rate(audio_np, frame_length=1024, hop_length=256)[0]
    flatness = librosa.feature.spectral_flatness(y=audio_np, n_fft=1024, hop_length=256)[0]
    onset_env = librosa.onset.onset_strength(y=audio_np, sr=SAMPLE_RATE)
    stft_mag = np.abs(librosa.stft(audio_np, n_fft=512, hop_length=256))

    eps = 1e-8
    rms_db = 20 * np.log10(rms + eps)
    dynamic_range_db = float(np.percentile(rms_db, 95) - np.percentile(rms_db, 5))
    clipping_ratio = float(np.mean(np.abs(audio_np) > 0.98))
    mean_zcr = float(np.mean(zcr))
    mean_flatness = float(np.mean(flatness))
    mean_onset = float(np.mean(onset_env)) if len(onset_env) > 0 else 0.0

    if stft_mag.shape[1] > 5:
        corrs = []
        for i in range(stft_mag.shape[1] - 1):
            a = stft_mag[:, i]
            b = stft_mag[:, i + 1]
            denom = (np.linalg.norm(a) * np.linalg.norm(b)) + 1e-8
            corrs.append(float(np.dot(a, b) / denom))
        mean_frame_corr = float(np.mean(corrs))
    else:
        mean_frame_corr = 0.0

    suspicion_flags = []

    if dynamic_range_db < MIN_DYNAMIC_RANGE_DB:
        suspicion_flags.append("low_dynamic_range")

    if clipping_ratio > MAX_CLIPPING_RATIO:
        suspicion_flags.append("heavy_clipping")

    if mean_zcr < MIN_ZERO_CROSS_RATE or mean_zcr > MAX_ZERO_CROSS_RATE:
        suspicion_flags.append("abnormal_zcr")

    if mean_flatness > MAX_SPECTRAL_FLATNESS:
        suspicion_flags.append("high_spectral_flatness")

    if mean_onset < MIN_ONSET_STRENGTH:
        suspicion_flags.append("weak_onset_pattern")

    if mean_frame_corr > MAX_REPEATED_FRAME_CORR:
        suspicion_flags.append("overstable_adjacent_frames")

    raw_suspicion = len(suspicion_flags) / 6.0
    liveness_score = round(max(0.0, 1.0 - raw_suspicion), 4)

    if len(suspicion_flags) >= 2:
        label = "PLAYBACK_SUSPECTED"
    else:
        label = "LIVE_LIKELY"

    return {
        "label": label,
        "score": liveness_score,
        "flags": suspicion_flags,
        "dynamic_range_db": round(dynamic_range_db, 2),
        "clipping_ratio": round(clipping_ratio, 4),
        "mean_zcr": round(mean_zcr, 4),
        "mean_flatness": round(mean_flatness, 4),
        "mean_onset": round(mean_onset, 4),
        "mean_frame_corr": round(mean_frame_corr, 4),
    }


def final_decision(prompt_ok, speaker_ok, liveness_live):
    if not prompt_ok:
        return "REJECTED - PROMPT MISMATCH"

    if not speaker_ok:
        return "REJECTED - UNKNOWN SPEAKER"

    if not liveness_live:
        return "REJECTED - PLAYBACK ATTACK SUSPECTED"

    return "ACCEPTED - PAVAN VERIFIED, LIVE SPEAKER"


# -------------------------------------------------
# MAIN
# -------------------------------------------------
print("\nInitializing Smart Speaker Verification + Liveness System...")

with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
    classifier = EncoderClassifier.from_hparams(
        source="speechbrain/spkrec-ecapa-voxceleb",
        savedir=ECAPA_CACHE_DIR,
        run_opts={"device": DEVICE}
    )

with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
    asr = pipeline(
        "automatic-speech-recognition",
        model="openai/whisper-medium",
        device=0 if DEVICE == "cuda" else -1
    )

print("System Ready.")

expected_prompt = random.choice(PROMPT_BANK)
print("\n" + "=" * 80)
print("RANDOM CHALLENGE PROMPT")
print("=" * 80)
print(expected_prompt)
print("=" * 80)
print("Please read the prompt clearly after speech is detected.")
print("Waiting for speech...")

# -------------------------------------------------
# CAPTURE AUDIO
# -------------------------------------------------
recording = capture_speech()
audio_np = prepare_audio(recording)

if audio_np is None or len(audio_np) < SAMPLE_RATE * 2:
    print("No valid speech captured or speech too short.")
    sys.exit()

# -------------------------------------------------
# SAVE INPUT AUDIO
# -------------------------------------------------
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
input_audio_file = os.path.abspath(os.path.join(RECORDINGS_DIR, f"input_{timestamp}.wav"))
sf.write(input_audio_file, audio_np, SAMPLE_RATE)

# -------------------------------------------------
# SPEECH TO TEXT
# -------------------------------------------------
try:
    transcription = asr(audio_np)
    spoken_text = transcription["text"].strip()
except Exception:
    spoken_text = "[ASR Failed]"

prompt_result = verify_prompt(expected_prompt, spoken_text)
prompt_match_score = prompt_result["similarity"]
prompt_word_coverage_score = prompt_result["coverage"]
prompt_ok = prompt_result["prompt_ok"]
missing_prompt_words = prompt_result["missing_words"]

# -------------------------------------------------
# SPEAKER VERIFICATION
# -------------------------------------------------
input_embedding = extract_embedding(classifier, audio_np)
speaker_metrics = compute_reference_scores(classifier, input_embedding, audio_np)

best_score = speaker_metrics["best_score"]
best_sample = speaker_metrics["best_sample"]
avg_similarity = speaker_metrics["avg_similarity"]

speaker_ok = (
    best_score >= BEST_SCORE_THRESHOLD and
    avg_similarity >= AVG_SCORE_THRESHOLD
)

detected_speaker = "PAVAN" if speaker_ok else "UNKNOWN"

# -------------------------------------------------
# LIVENESS / PLAYBACK SUSPICION
# -------------------------------------------------
liveness = estimate_liveness(audio_np)
liveness_live = liveness["label"] == "LIVE_LIKELY"

# -------------------------------------------------
# FINAL DECISION
# -------------------------------------------------
decision = final_decision(
    prompt_ok=prompt_ok,
    speaker_ok=speaker_ok,
    liveness_live=liveness_live
)

# -------------------------------------------------
# OUTPUT
# -------------------------------------------------
print("\n" + "=" * 80)
print("SMART SPEAKER AUTHENTICATION RESULT")
print("=" * 80)
print("Expected Prompt       :", expected_prompt)
print("Spoken Text           :", spoken_text)
print("Prompt Match Score    :", prompt_match_score)
print("Prompt Word Coverage  :", prompt_word_coverage_score)
print("Missing Prompt Words  :", ", ".join(missing_prompt_words) if missing_prompt_words else "None")
print("Prompt Verification   :", "PASS" if prompt_ok else "FAIL")
print()
print("Detected Speaker      :", detected_speaker)
print("Best Matching Sample  :", best_sample)
print("Best Similarity Score :", round(best_score, 4))
print("Average Similarity    :", round(avg_similarity, 4))
print()
print("Liveness Label        :", liveness["label"])
print("Liveness Score        :", liveness["score"])
print("Liveness Flags        :", ", ".join(liveness["flags"]) if liveness["flags"] else "None")
print()
print("Pitch Difference      :", speaker_metrics["pitch_diff"])
print("Spectral Centroid Diff:", speaker_metrics["centroid_diff"])
print("Bandwidth Difference  :", speaker_metrics["bandwidth_diff"])
print("Energy Difference     :", speaker_metrics["energy_diff"])
print("MFCC Distance         :", speaker_metrics["mfcc_diff"])
print()
print("FINAL DECISION        :", decision)
print("=" * 80)

# -------------------------------------------------
# SAVE RESULTS TO WORKBOOK
# -------------------------------------------------
excel_file = os.path.join(LOGS_DIR, "speaker_verification_log.xlsx")

safe_close_libreoffice()
repair_existing_audio_links(excel_file, BASE_DIR)

sheet_name = get_next_run_sheet_name(excel_file)

if best_sample:
    match_audio_file = os.path.abspath(os.path.join(PAVAN_RECORDINGS_DIR, best_sample))
else:
    match_audio_file = ""

data = {
    "Timestamp": [datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    "Expected_Prompt": [expected_prompt],
    "Spoken_Text": [spoken_text],
    "Prompt_Match": [prompt_match_score],
    "Prompt_Word_Coverage": [prompt_word_coverage_score],
    "Missing_Prompt_Words": [", ".join(missing_prompt_words) if missing_prompt_words else "None"],
    "Detected_Speaker": [detected_speaker],
    "Input_Audio": [input_audio_file],
    "Matched_Audio": [match_audio_file],
    "Best_Matching_Sample": [best_sample],
    "Best_Similarity": [round(best_score, 4)],
    "Average_Similarity": [round(avg_similarity, 4)],
    "Pitch_Difference": [speaker_metrics["pitch_diff"]],
    "Spectral_Centroid_Diff": [speaker_metrics["centroid_diff"]],
    "Bandwidth_Difference": [speaker_metrics["bandwidth_diff"]],
    "Energy_Difference": [speaker_metrics["energy_diff"]],
    "MFCC_Distance": [speaker_metrics["mfcc_diff"]],
    "Liveness_Label": [liveness["label"]],
    "Liveness_Score": [liveness["score"]],
    "Decision": [decision]
}

df = pd.DataFrame(data)

if os.path.exists(excel_file):
    with pd.ExcelWriter(excel_file, engine="openpyxl", mode="a", if_sheet_exists="new") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
else:
    with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

workbook = load_workbook(excel_file)
worksheet = workbook[sheet_name]
format_sheet(worksheet, max_width=42)

input_col = None
match_col = None

for col in range(1, worksheet.max_column + 1):
    header = worksheet.cell(row=1, column=col).value
    if header == "Input_Audio":
        input_col = col
    elif header == "Matched_Audio":
        match_col = col

if input_col is not None and input_audio_file and os.path.exists(input_audio_file):
    cell = worksheet.cell(row=2, column=input_col)
    cell.value = "▶ Play Input"
    cell.hyperlink = os.path.abspath(input_audio_file)
    cell.style = "Hyperlink"

if match_col is not None and match_audio_file and os.path.exists(match_audio_file):
    cell = worksheet.cell(row=2, column=match_col)
    cell.value = "▶ Play Match"
    cell.hyperlink = os.path.abspath(match_audio_file)
    cell.style = "Hyperlink"

workbook.save(excel_file)
time.sleep(1)

repair_existing_audio_links(excel_file, BASE_DIR)
rebuild_master_log(excel_file)

print("\nLog saved to:", excel_file)